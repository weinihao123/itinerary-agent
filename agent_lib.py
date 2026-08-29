# -*- coding: utf-8 -*-
"""行程记录与分析智能体 核心库

职责：
- CSV 单一数据源的读写
- 派生天数计算（含首尾口径）
- 讲师层级逐期 +1（同项目）
- 编辑 / 软删除 / 物理删除
- Excel 导出（明细 + 汇总）
- 看板 HTML 生成（Chart.js）
- 口语文本解析为草稿（尽力而为，低置信字段交由对话层追问）
"""
import csv
import json
import os
import re
from datetime import date, datetime, timedelta

BASE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE, "data")
CSV_PATH = os.path.join(DATA_DIR, "trips.csv")

COLUMNS = [
    "记录ID", "项目名称", "运营商", "执行省份", "执行地市", "项目类型",
    "项目执行开始", "项目执行结束", "出发日期", "返程日期",
    "项目执行天数", "项目总计花费时间", "讲师层级",
    "录入时间", "最后修改", "作废标记",
]

PROJECT_TYPES = ["AI", "专线", "云", "ICT", "商客", "其他"]
OPERATORS = ["移动", "联通", "电信", "广电"]

PROVINCES = [
    "北京", "天津", "河北", "山西", "内蒙古", "辽宁", "吉林", "黑龙江", "上海", "江苏",
    "浙江", "安徽", "福建", "江西", "山东", "河南", "湖北", "湖南", "广东", "广西", "海南",
    "重庆", "四川", "贵州", "云南", "西藏", "陕西", "甘肃", "青海", "宁夏", "新疆",
    "台湾", "香港", "澳门",
]

# 主要城市 -> 省份（用于口语解析时反推省份）
CITIES = {
    "郑州": "河南", "洛阳": "河南", "开封": "河南", "新乡": "河南", "南阳": "河南",
    "石家庄": "河北", "唐山": "河北", "太原": "山西", "呼和浩特": "内蒙古",
    "沈阳": "辽宁", "大连": "辽宁", "长春": "吉林", "哈尔滨": "黑龙江",
    "南京": "江苏", "苏州": "江苏", "无锡": "江苏", "徐州": "江苏", "扬州": "江苏",
    "杭州": "浙江", "宁波": "浙江", "温州": "浙江", "合肥": "安徽", "福州": "福建",
    "厦门": "福建", "南昌": "江西", "济南": "山东", "青岛": "山东", "烟台": "山东",
    "潍坊": "山东", "武汉": "湖北", "长沙": "湖南", "广州": "广东", "深圳": "广东",
    "东莞": "广东", "珠海": "广东", "南宁": "广西", "海口": "海南", "成都": "四川",
    "重庆": "重庆", "贵阳": "贵州", "昆明": "云南", "拉萨": "西藏", "西安": "陕西",
    "兰州": "甘肃", "西宁": "青海", "银川": "宁夏", "乌鲁木齐": "新疆",
    "北京": "北京", "上海": "上海", "天津": "天津",
}


# ---------------- 存储 ----------------
def ensure_store():
    os.makedirs(DATA_DIR, exist_ok=True)
    if not os.path.exists(CSV_PATH):
        with open(CSV_PATH, "w", encoding="utf-8", newline="") as f:
            csv.writer(f).writerow(COLUMNS)


def load_rows():
    ensure_store()
    with open(CSV_PATH, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def save_rows(rows):
    with open(CSV_PATH, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in COLUMNS})


def active_rows(rows=None):
    rows = rows if rows is not None else load_rows()
    return [r for r in rows if str(r.get("作废标记", "0")).strip() not in ("1", "true", "True")]


def next_id(rows=None):
    rows = rows if rows is not None else load_rows()
    ids = [int(r["记录ID"]) for r in rows if str(r.get("记录ID", "")).strip().isdigit()]
    return (max(ids) + 1) if ids else 1


# ---------------- 日期工具 ----------------
def _this_year():
    return date.today().year


def normalize_date(s):
    if not s:
        return ""
    s = s.strip()
    m = re.match(r"^(\d{4})[-/](\d{1,2})[-/](\d{1,2})$", s)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    m = re.match(r"^(\d{1,2})[-/](\d{1,2})$", s)
    if m:
        return f"{_this_year():04d}-{int(m.group(1)):02d}-{int(m.group(2)):02d}"
    # 支持 "号" 与 "日"
    m = re.search(r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*[号日]", s)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    m = re.search(r"(\d{1,2})\s*月\s*(\d{1,2})\s*[号日]", s)
    if m:
        return f"{_this_year():04d}-{int(m.group(1)):02d}-{int(m.group(2)):02d}"
    return ""


def _datediff_inc(a, b):
    """含首尾天数：b - a + 1"""
    da = datetime.strptime(normalize_date(a), "%Y-%m-%d").date()
    db = datetime.strptime(normalize_date(b), "%Y-%m-%d").date()
    return (db - da).days + 1


def compute_derived(fields):
    out = {}
    if normalize_date(fields.get("项目执行开始", "")) and normalize_date(fields.get("项目执行结束", "")):
        out["项目执行天数"] = str(_datediff_inc(fields["项目执行开始"], fields["项目执行结束"]))
    if normalize_date(fields.get("出发日期", "")) and normalize_date(fields.get("返程日期", "")):
        out["项目总计花费时间"] = str(_datediff_inc(fields["出发日期"], fields["返程日期"]))
    return out


# ---------------- 讲师层级 ----------------
def compute_level(rows, project_name, override=None):
    """同项目已确认（活跃）记录的最大层级 +1；首期默认 1。override 用于用户改写首期。"""
    levels = []
    for r in active_rows(rows):
        if r.get("项目名称", "").strip() == project_name.strip():
            try:
                levels.append(int(float(r.get("讲师层级", "0"))))
            except (ValueError, TypeError):
                pass
    base = (max(levels) + 1) if levels else (override if override is not None else 1)
    return base


# ---------------- 口语解析 ----------------
def _expand_date_ranges(text):
    """把 '9月3号和4号' 展开为 '9月3号到9月4号'，便于通用区间识别"""
    text = re.sub(
        r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*[号日]\s*[,，、和与至到～\-]\s*(\d{1,2})\s*[号日]",
        r"\1年\2月\3号到\1年\2月\4号",
        text,
    )
    text = re.sub(
        r"(\d{1,2})\s*月\s*(\d{1,2})\s*[号日]\s*[,，、和与至到～\-]\s*(\d{1,2})\s*[号日]",
        r"\1月\2号到\1月\3号",
        text,
    )
    return text


def _find_dates(text):
    """返回 [(iso, start_idx, end_idx), ...] 按出现顺序；调用方应已展开日期连写"""
    pats = [
        r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*[号日]",
        r"(\d{1,2})\s*月\s*(\d{1,2})\s*[号日]",
        r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})",
        r"(\d{1,2})[-/](\d{1,2})",
    ]
    results = []
    for pat in pats:
        for m in re.finditer(pat, text):
            iso = normalize_date(m.group(0))
            if iso:
                results.append((iso, m.start(), m.end()))
    # 去重重叠（长模式优先），按起始排序
    results.sort(key=lambda x: x[1])
    cleaned = []
    for iso, s, e in results:
        if cleaned and s < cleaned[-1][2]:
            continue
        cleaned.append((iso, s, e))
    return cleaned


def _extract_project_name(raw, found):
    """从口语中抽取或构造项目名称。若用户明确说“XX项目”则优先抽取；否则按城市+运营商+类型构造。"""
    name = ""
    # 仅当文本中明确出现“项目”时才尝试抽取；否则直接构造，避免把日期、口语套语也抓进来
    if "项目" in raw:
        m = re.search(r"(.+?项目)", raw)
        if m:
            name = m.group(1).strip()
            name = re.sub(r"^(我要|我准备|我计划|我将|我于|我在|我要到|我到|我去)\s*", "", name)
            name = re.sub(r"^\d{4}年\d{1,2}月\d{1,2}[号日](\s*[到至去])?\s*", "", name)
            name = re.sub(r"^\d{1,2}月\d{1,2}[号日](\s*[到至去])?\s*", "", name)
            name = re.sub(r"^[到至去]\s*", "", name)
            name = name.replace("执行", "")
            name = re.sub(r"的?项目$", "项目", name)
    # 构造默认值
    if not name or len(name) < 2 or re.search(r"\d|[号日月年]|我要|到", name):
        parts = [found.get("执行地市", ""), found.get("运营商", ""), found.get("项目类型", "")]
        activity = "项目"
        if "培训" in raw:
            activity = "培训"
        elif "训战" in raw:
            activity = "训战"
        elif "课程" in raw:
            activity = "课程"
        elif "工作坊" in raw:
            activity = "工作坊"
        elif "讲座" in raw:
            activity = "讲座"
        name = "".join(p for p in parts if p) + activity
    return name


def _parse_level(text):
    """识别文本中明确的当前讲师层级，例如'目前等级是金牌23级'。返回 int 或 None。"""

    def to_int(s):
        if s.isdigit():
            return int(s)
        # 简单中文数字（仅个位数）
        m = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5,
             "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}
        if s in m:
            return m[s]
        return None

    # 含“等级/职级/级别”：目前等级呢是金牌23级、现在职级是23级
    m = re.search(r"(?:目前|当前|现在|这[一]?期)?\s*(?:等级|职级|级别)\D*?([0-9一二三四五六七八九十]+)\s*级", text)
    if m:
        return to_int(m.group(1))
    # 不含“等级/职级/级别”：目前金牌23级、这一期23级
    m = re.search(r"(?:目前|当前|现在|这[一]?期)[是为]?[了]?\s*(?:金牌|银牌|铜牌|钻石|白金)?\s*([0-9一二三四五六七八九十]+)\s*级", text)
    if m:
        return to_int(m.group(1))
    return None


def parse_oral_text(text):
    """尽力解析口语文本，返回 {'found': {...}, 'missing': [...], 'level_hint': int|None}"""
    found = {}
    raw = _expand_date_ranges(text)

    # 运营商
    for op in OPERATORS:
        if op in raw:
            found["运营商"] = op
            break

    # 项目类型
    for t in PROJECT_TYPES:
        if t in raw:
            found["项目类型"] = t
            break

    # 省份 / 地市（项目名称依赖省份/地市，先解析）
    prov = None
    city = None
    for p in PROVINCES:
        if p in raw:
            prov = p
            break
    for c, pr in CITIES.items():
        if c in raw:
            city = c
            prov = prov or pr
            break
    if prov:
        found["执行省份"] = prov
    if city:
        found["执行地市"] = city

    # 项目名称
    found["项目名称"] = _extract_project_name(raw, found)

    # 日期
    dates = _find_dates(raw)
    used = set()
    # 1) 识别区间（A到/至B）作为执行时间
    for i in range(len(dates) - 1):
        a, as_, ae = dates[i]
        b, bs, be = dates[i + 1]
        mid = raw[ae:bs]
        if "到" in mid or "至" in mid:
            found["项目执行开始"] = a
            found["项目执行结束"] = b
            used.add(i)
            used.add(i + 1)
            break

    # 2) 出发 / 返程 关键词邻近（前后各看4字）
    remain = [(j, d) for j, d in enumerate(dates) if j not in used]
    for j, (iso, s, e) in remain:
        ctx = raw[max(0, s - 4):s] + raw[e:e + 4]
        if "出发日期" not in found and any(k in ctx for k in ("出发", "走", "去", "启程")):
            found["出发日期"] = iso
            used.add(j)
        elif "返程日期" not in found and any(k in ctx for k in ("返", "回", "回程", "归")):
            found["返程日期"] = iso
            used.add(j)

    # 3) 残余日期按序填补出发/返程
    free_idx = [j for j in range(len(dates)) if j not in used]
    if "出发日期" not in found and free_idx:
        found["出发日期"] = dates[free_idx[0]][0]
        used.add(free_idx[0])
    if "返程日期" not in found and len(free_idx) > 1:
        found["返程日期"] = dates[free_idx[1]][0]
        used.add(free_idx[1])

    # 缺失项
    required = ["项目名称", "运营商", "执行省份", "执行地市", "项目类型",
                "项目执行开始", "项目执行结束", "出发日期", "返程日期"]
    missing = [k for k in required if not found.get(k)]
    level_hint = _parse_level(text)
    return {"found": found, "missing": missing, "level_hint": level_hint}


# ---------------- 写操作 ----------------
def add_trip(fields, level_override=None):
    """fields 为录入字段字典（不含派生/层级/元数据）。返回新建记录。"""
    rows = load_rows()
    rec = {c: "" for c in COLUMNS}
    for k in ["项目名称", "运营商", "执行省份", "执行地市", "项目类型",
              "项目执行开始", "项目执行结束", "出发日期", "返程日期"]:
        rec[k] = fields.get(k, "").strip()
    rec.update(compute_derived(rec))
    proj = rec["项目名称"]
    rec["讲师层级"] = str(compute_level(rows, proj, override=level_override))
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    rec["记录ID"] = str(next_id(rows))
    rec["录入时间"] = now
    rec["最后修改"] = now
    rec["作废标记"] = "0"
    rows.append(rec)
    save_rows(rows)
    return rec


def edit_trip(record_id, updates):
    rows = load_rows()
    for r in rows:
        if r.get("记录ID") == str(record_id):
            for k, v in updates.items():
                if k in COLUMNS:
                    r[k] = v
            r.update(compute_derived(r))
            r["最后修改"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            save_rows(rows)
            return r
    return None


def soft_delete(record_id):
    return edit_trip(record_id, {"作废标记": "1"})


def hard_delete(record_id):
    rows = load_rows()
    rows = [r for r in rows if r.get("记录ID") != str(record_id)]
    save_rows(rows)
    return True


# ---------------- Excel 导出 ----------------
def export_excel(path=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    if path is None:
        path = os.path.join(BASE, "trips.xlsx")
    rows = active_rows()
    wb = Workbook()

    # 明细
    ws = wb.active
    ws.title = "明细"
    ws.append(COLUMNS)
    head_fill = PatternFill("solid", fgColor="185FA5")
    for c in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([r.get(c, "") for c in COLUMNS])
    ws.freeze_panes = "A2"

    # 汇总
    ws2 = wb.create_sheet("汇总")
    ws2.append(["维度", "分类", "出行次数", "执行天数合计", "花费时间合计"])
    for c in range(1, 6):
        cell = ws2.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill

    def _num(v):
        try:
            return int(float(v))
        except (ValueError, TypeError):
            return 0

    def _agg(key):
        agg = {}
        for r in rows:
            if key == "年份":
                k = (r.get("出发日期") or "")[:4] or "(空)"
            else:
                k = r.get(key, "(空)")
            a = agg.setdefault(k, [0, 0, 0])
            a[0] += 1
            a[1] += _num(r.get("项目执行天数"))
            a[2] += _num(r.get("项目总计花费时间"))
        for k in sorted(agg):
            ws2.append([key, k, agg[k][0], agg[k][1], agg[k][2]])

    _agg("执行省份")
    _agg("项目类型")
    _agg("运营商")
    _agg("年份")

    wb.save(path)
    return path


# ---------------- 看板 ----------------
def _bucket(iso, dim):
    d = datetime.strptime(normalize_date(iso), "%Y-%m-%d").date()
    if dim == "周":
        y, w, _ = d.isocalendar()
        return f"{y}-W{w:02d}"
    if dim == "月":
        return f"{d.year}-{d.month:02d}"
    return str(d.year)


def get_dashboard_data():
    """返回看板所需的结构化记录列表（不含 HTML）"""
    rows = active_rows()
    data = []
    for r in rows:
        data.append({
            "id": r.get("记录ID"),
            "project": r.get("项目名称"),
            "operator": r.get("运营商"),
            "province": r.get("执行省份"),
            "city": r.get("执行地市"),
            "type": r.get("项目类型"),
            "exec_start": r.get("项目执行开始"),
            "exec_end": r.get("项目执行结束"),
            "depart": r.get("出发日期"),
            "return": r.get("返程日期"),
            "exec_days": _to_int(r.get("项目执行天数")),
            "total_days": _to_int(r.get("项目总计花费时间")),
            "level": _to_int(r.get("讲师层级")),
            "year": (r.get("出发日期") or "")[:4],
        })
    return data


def build_dashboard(path=None):
    if path is None:
        path = os.path.join(BASE, "dashboard.html")
    data = get_dashboard_data()
    html = _DASHBOARD_TPL.replace("__DATA__", json.dumps(data, ensure_ascii=False))
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    return path


def _to_int(v):
    try:
        return int(float(v))
    except (ValueError, TypeError):
        return 0


_DASHBOARD_TPL = r"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>行程分析看板</title>
<script src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.1/chart.umd.min.js"></script>
<style>
  body{font-family:-apple-system,"PingFang SC","Microsoft YaHei",sans-serif;margin:0;background:#f8fafc;color:#0f172a}
  header{padding:16px 24px;background:#0c447c;color:#fff}
  header h1{font-size:18px;margin:0}
  .wrap{padding:16px 24px}
  .filters{display:flex;flex-wrap:wrap;gap:12px;background:#fff;padding:14px;border-radius:12px;box-shadow:0 1px 3px rgba(0,0,0,.08)}
  .filters label{font-size:13px;color:#475569;display:flex;flex-direction:column;gap:4px}
  select{padding:6px 10px;border:1px solid #cbd5e1;border-radius:8px;font-size:13px;min-width:120px}
  .grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(320px,1fr));gap:16px;margin-top:16px}
  .card{background:#fff;border-radius:12px;padding:14px;box-shadow:0 1px 3px rgba(0,0,0,.08)}
  .card h3{margin:0 0 10px;font-size:14px;color:#1e293b}
  .empty{padding:40px;text-align:center;color:#94a3b8}
</style>
</head>
<body>
<header><h1>行程记录与分析看板</h1></header>
<div class="wrap">
  <div class="filters">
    <label>时间维度<select id="dim"><option value="周">周</option><option value="月">月</option><option value="年">年</option></select></label>
    <label>省份<select id="province"></select></label>
    <label>地市<select id="city"></select></label>
    <label>项目类型<select id="type"></select></label>
    <label>运营商<select id="operator"></select></label>
  </div>
  <div id="charts" class="grid">
    <div class="card"><h3>出行次数</h3><canvas id="c1"></canvas></div>
    <div class="card"><h3>执行天数合计</h3><canvas id="c2"></canvas></div>
    <div class="card"><h3>花费时间合计</h3><canvas id="c3"></canvas></div>
  </div>
  <div id="empty" class="empty" style="display:none">当前筛选无数据，请调整条件。</div>
</div>
<script>
const DATA = __DATA__;
const charts = {};
function uniq(arr){return [...new Set(arr)].filter(Boolean).sort();}
function fillSelect(id, vals){const s=document.getElementById(id);const cur=s.value;s.innerHTML='<option value="">全部</option>'+vals.map(v=>`<option>${v}</option>`).join('');if(vals.includes(cur))s.value=cur;}
function bucket(iso,dim){if(!iso)return'';const d=new Date(iso);if(dim==='周'){const t=new Date(d.getFullYear(),0,1);const w=Math.ceil((((d-t)/86400000)+t.getDay()+1)/7);return d.getFullYear()+'-W'+String(w).padStart(2,'0');}if(dim==='月')return d.getFullYear()+'-'+String(d.getMonth()+1).padStart(2,'0');return String(d.getFullYear());}
function render(){
  const dim=document.getElementById('dim').value;
  const f={province:document.getElementById('province').value,city:document.getElementById('city').value,type:document.getElementById('type').value,operator:document.getElementById('operator').value};
  const rows=DATA.filter(r=>(!f.province||r.province===f.province)&&(!f.city||r.city===f.city)&&(!f.type||r.type===f.type)&&(!f.operator||r.operator===f.operator));
  const labels=uniq(rows.map(r=>bucket(r.depart,dim)));
  const cnt=labels.map(l=>rows.filter(r=>bucket(r.depart,dim)===l).length);
  const ed=labels.map(l=>rows.filter(r=>bucket(r.depart,dim)===l).reduce((a,r)=>a+r.exec_days,0));
  const td=labels.map(l=>rows.filter(r=>bucket(r.depart,dim)===l).reduce((a,r)=>a+r.total_days,0));
  const show=rows.length>0;
  document.getElementById('charts').style.display=show?'grid':'none';
  document.getElementById('empty').style.display=show?'none':'block';
  if(!show)return;
  const mk=(id,type,data,label,color)=>{if(charts[id])charts[id].destroy();const ctx=document.getElementById(id);charts[id]=new Chart(ctx,{type,data:{labels,datasets:[{label,data,backgroundColor:color,borderColor:color,tension:.2,fill:type==='line'}]},options:{responsive:true,plugins:{legend:{display:true}}}});};
  mk('c1','bar',cnt,'次数','#185FA5');
  mk('c2','bar',ed,'执行天数','#3B6D11');
  mk('c3','line',td,'花费时间','#BA7517');
}
function init(){fillSelect('province',uniq(DATA.map(r=>r.province)));fillSelect('city',uniq(DATA.map(r=>r.city)));fillSelect('type',uniq(DATA.map(r=>r.type)));fillSelect('operator',uniq(DATA.map(r=>r.operator)));['dim','province','city','type','operator'].forEach(id=>document.getElementById(id).addEventListener('change',render));render();}
init();
</script>
</body>
</html>
"""
